import logging
import re
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from docx import Document
from .schema import NormalizedData

logger = logging.getLogger(__name__)

class FormAdapter:
    @staticmethod
    def get_concept_labels(form_type: str) -> Dict[str, List[str]]:
        if form_type == "power_of_attorney":
            return {
                "people.principal.full_name": ["alulírott", "én,", "név:"],
                "people.principal.address": ["lakóhely:", "lakcím:"],
                "people.principal.id_number": ["személyazonosító igazolványának száma:", "szem. ig. szám.:"],
                "people.principal.mother_name": ["anyja neve:"],
                "people.principal.birth_date": ["születési hely, idő:"],
                "people.agent.full_name": ["meghatalmazom"],
                "people.agent.address": ["lakcím:"],
                "people.agent.id_number": ["személyazonosító igazolványának száma:"],
                "people.agent.mother_name": ["anyja neve:"],
                "people.agent.birth_date": ["születési hely, idő:"],
                "case.authority_name": ["(eljáró hatóság megnevezése)", "hatóság megnevezése"],
                "case.case_type": ["(ügy típusának megjelölése)", "ügy típusának megjelölése"],
                "case.city": ["kelt:"],
                "case.date": ["év", "hónap", "nap"],
                "people.witness_1.full_name": ["név:"],
                "people.witness_1.address": ["lakcím:"],
                "people.witness_1.id_number": ["szem. ig. szám.:"],
                "people.witness_2.full_name": ["név:"],
                "people.witness_2.address": ["lakcím:"],
                "people.witness_2.id_number": ["szem. ig. szám.:"],
            }
        elif form_type == "family_tax_benefit":
            return {
                "tax.submission_year": ["benyújtásának éve:"],
                "people.taxpayer.full_name": ["magánszemély neve:"],
                "people.taxpayer.tax_id": ["adóazonosító jele:"],
                "tax.hungary_only_claim": ["Magyarországon érvényesíteni", "csak Magyarországon"],
                "tax.disable_contribution_discount": ["nem kérem a családi járulékkedvezmény", "mellőzése"],
                "tax.discount_amount_huf": ["forint összegben kívánom"],
                "tax.beneficiary_count": ["fő kedvezményezett eltartott után"],
                "people.partner.full_name": ["neve "],
                "people.partner.tax_id": ["adóazonosító jele:"],
                "tax.partner_employer_name": ["kifizetője megnevezése:"],
                "tax.partner_employer_tax_number": ["adószáma:"],
                "tax.employer_name": ["kifizető megnevezése:"],
                "tax.employer_tax_number": ["adószáma:"],
                "tax.dependents.name": ["név"],
                "tax.dependents.tax_id": ["adóazonosító jel"],
                "tax.dependents.em_code": ["em*"],
                "tax.dependents.jj_code": ["jj**"],
                "tax.dependents.change_date": ["változás időpontja"],
            }
        return {}

class FormFiller:
    def __init__(self):
        self.current_section = None
        self.current_dep_idx = 0
        self.runs_to_replace = []

    def detect_type(self, doc_path: Path) -> str:
        try:
            doc = Document(doc_path)
            ft = "\n".join([p.text for p in doc.paragraphs]).lower()
            if "meghatalmazás" in ft: return "power_of_attorney"
            if "családi" in ft or "nyilatkozat" in ft: return "family_tax_benefit"
            return "unknown"
        except: return "unknown"

    def fill(self, template_path: Path, data: NormalizedData, output_path: Path):
        form_type = self.detect_type(template_path)
        labels_map = FormAdapter.get_concept_labels(form_type)
        
        if form_type == "family_tax_benefit":
            if data.tax.beneficiary_count and data.tax.discount_amount_huf:
                data.tax.discount_amount_huf = ""
                
        doc = Document(template_path)
        self.current_section = None
        self.current_dep_idx = 0
        
        paragraphs = []
        for p in doc.paragraphs: paragraphs.append(p)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs: paragraphs.append(p)
        
        for p in paragraphs:
            self._fill_p(p, labels_map, data)
        
        # APPLY CHANGES
        for p, text in self.runs_to_replace:
             if "Magyarországon érvényesíteni" in text:
                  print(">>> FLUSHING P14 WITH TEXT:", repr(text))
             p.text = text
             
        doc.save(output_path)
        logger.info(f"Saved filled DOCX to {output_path}")

    def _fill_p(self, paragraph, labels_map, data):
        text = paragraph.text
        if not text.strip(): return
        t_low = text.lower()
        original_text = text
        
        # Unify disjointed tax number placeholders
        if "⎕⎕⎕⎕⎕⎕⎕⎕—⎕—⎕⎕" in text:
             text = text.replace("⎕⎕⎕⎕⎕⎕⎕⎕—⎕—⎕⎕", "⎕⎕⎕⎕⎕⎕⎕⎕⎕⎕⎕")
        if "tanú 1.:" in t_low: self.current_section = "witness_1"
        elif "tanú 2.:" in t_low: self.current_section = "witness_2"
        elif "meghatalmazom" in t_low: self.current_section = "agent"
        elif "én," in t_low or "alulírott" in t_low: self.current_section = "principal"
        elif "nyilatkozó magánszemély neve" in t_low: self.current_section = "taxpayer"
        elif "kelt:" in t_low: self.current_section = "signature"
        elif "jogosult házastársa/élettársa" in t_low: self.current_section = "partner"
        elif "magánszemély munkáltatójaként" in t_low: self.current_section = "employer"
        elif "eltartottak adatai" in t_low or "adóazonosító jel \tnév" in t_low:
             self.current_section = "tax_details"

        # Unify disjointed tax number placeholders
        if "⎕⎕⎕⎕⎕⎕⎕⎕—⎕—⎕⎕" in text:
             text = text.replace("⎕⎕⎕⎕⎕⎕⎕⎕—⎕—⎕⎕", "⎕⎕⎕⎕⎕⎕⎕⎕⎕⎕⎕")

        person_prefix = f"people.{self.current_section}." if self.current_section and self.current_section not in ["tax_details", "signature", "employer"] else ""
        P_RE = r"([\.…_\-]{2,}|⎕+)"

        if text.strip().startswith(('.', '…')) and self.current_section == "agent":
             val = self._get_data_value(data, "people.agent.full_name")
             if val: 
                  text = re.sub(P_RE, self._get_formatted_val(val, "……………….", "people.agent.full_name"), text, count=1)
                  self.runs_to_replace.append((p, text))
                  return
                  
        if "adóazonosító jele:" in t_low and "\t⎕" in text:
             val = self._get_data_value(data, "tax.is_modified")
             text = text.replace("\t⎕", "\t" + self._get_formatted_val(val, "⎕", "tax.is_modified"), 1)
             
        if "egyedül ⎕" in text and "közösen ⎕" in text:
             val = self._get_data_value(data, "tax.joint_claim")
             if val is not None and str(val).strip() != "" and str(val).strip() != "None":
                  if val: text = text.replace("egyedül ⎕", "egyedül  ⎕").replace("közösen ⎕", "közösen  ☒")
                  else: text = text.replace("egyedül ⎕", "egyedül  ☒").replace("közösen ⎕", "közösen  ⎕")

        filled_any = False; used_ids = set(); dep_row_filled = False
        changed_in_pass = True
        
        while changed_in_pass:
             changed_in_pass = False; scores = []
             for ms in re.finditer(P_RE, text):
                  ps, pe = ms.span()
                  
                  # POSITIVE DATE LINE
             if not changed_in_pass: break

        ms = list(re.finditer(P_RE, text))
        if ms and self.current_section == "signature" and "év" in t_low and len(ms) >= 3:
             vals = [self._get_data_value(data, "case.city"), self._get_data_value(data, "case.date", "év"),
                     self._get_data_value(data, "case.date", "hónap"), self._get_data_value(data, "case.date", "nap")]
             idx = 0
             def date_sub(m):
                  nonlocal idx; v = vals[idx] if idx < len(vals) else ""; idx += 1
                  return f" {v} " if v else m.group(1)
             text = re.sub(P_RE, date_sub, text, count=4); filled_any = True

        # POSITIONAL DEPENDENT ROW
        elif ms and self.current_section == "tax_details" and ("⎕⎕⎕⎕⎕⎕" in original_text) and len(ms) >= 3:
             tid = self._get_data_value(data, "tax.dependents.tax_id")
             nm = self._get_data_value(data, "tax.dependents.name")
             em = self._get_data_value(data, "tax.dependents.em_code")
             jj = self._get_data_value(data, "tax.dependents.jj_code")
             dy = self._get_data_value(data, "tax.dependents.change_date", "év")
             dm = self._get_data_value(data, "tax.dependents.change_date", "hónap")
             dd = self._get_data_value(data, "tax.dependents.change_date", "nap")
             
             if tid or nm or em or jj:
                  dt_keys = ["tax.dependents.tax_id", "tax.dependents.name", "tax.dependents.em_code", "tax.dependents.jj_code", "tax.dependents.change_date", "tax.dependents.change_date", "tax.dependents.change_date"]
                  dt_idx = 0
                  def dep_sub(m):
                      nonlocal dt_idx
                      path = dt_keys[dt_idx]
                      v = self._get_data_value(data, path)
                      if path == "tax.dependents.change_date":
                           parts = re.split(r'[\.\-\s]+', str(v)) if v else []
                           if dt_idx == 4: # First date placeholder (YY)
                                v = parts[0][-2:] if len(parts) > 0 and len(parts[0]) >= 2 else ""
                           elif dt_idx == 5: # Second (MM)
                                v = parts[1] if len(parts) > 1 else ""
                           elif dt_idx == 6: # Third (DD)
                                v = parts[2] if len(parts) > 2 else ""
                      dt_idx += 1
                      return self._get_formatted_val(v, m.group(1), path)
                  text = re.sub(P_RE, dep_sub, text, count=7)
                  filled_any = True; dep_row_filled = True
             else:
                  # Clear out the unpopulated dependent row placeholders for visual cleanliness
                  text = re.sub(r"⎕", " ", text)
                  text = re.sub(r"[\.…_\-]{2,}", lambda m: " " * len(m.group(0)), text)
                  filled_any = True; dep_row_filled = True

        else:
             # Scored fallback
             while True:
                  changed_in_pass = False
                  curr_ms = list(re.finditer(P_RE, text))
                  if not curr_ms: break
                  scores = []
                  for m in curr_ms:
                       ph = m.group(1); ps, pe = m.start(), m.end()
                       if any(c not in "⎕.…_- " for c in ph): continue 
                       for path, lbls in labels_map.items():
                            if person_prefix and path.startswith("people.") and not path.startswith(person_prefix): continue
                            if self.current_section == "partner" and path.startswith("tax.employer_"): continue
                            if self.current_section == "employer" and path.startswith("tax.partner_employer_"): continue
                            if "tax.dependents" in path: continue
                            for lbl in lbls:
                                 l_id = f"{path}|{lbl}"
                                 if l_id in used_ids: continue
                                 for lm in re.finditer(re.escape(lbl), text, re.IGNORECASE):
                                      d = 999
                                      if lm.end() <= ps: d = ps - lm.end()
                                      elif lm.start() >= pe: d = lm.start() - pe
                                      else: continue
                                      
                                      max_d = 150 if "⎕" in ph else 60
                                      if d < max_d:
                                           b = text[min(lm.end(), pe):max(lm.start(), ps)]
                                           if not re.search(r"[\.…_]{2,}", b):
                                                scores.append({"score": d, "m": m, "path": path, "label": lbl, "id": l_id})
                  
                  scores.sort(key=lambda x: x["score"])
                  if scores:
                       best = scores[0]; val = self._get_data_value(data, best["path"], context_label=best["label"])
                       if val is not None and str(val).strip():
                            text = text[:best["m"].start()] + self._get_formatted_val(val, best["m"].group(1), best["path"]) + text[best["m"].end():]
                            used_ids.add(best["id"]); changed_in_pass = True; filled_any = True
                  if not changed_in_pass: break

        if dep_row_filled:
             self.current_dep_idx += 1
        
        if self.current_section == "signature" and "\t" in text:
             pv, av = self._get_data_value(data, "people.principal.full_name"), self._get_data_value(data, "people.agent.full_name")
             if pv or av:
                  parts = text.split("\t"); nps = []
                  for p in parts:
                       if re.search(r"[\.…]{3,}", p):
                            if pv and not any(str(pv) in s for s in nps): nps.append(f" {pv} ")
                            elif av and not any(str(av) in s for s in nps): nps.append(f" {av} ")
                            else: nps.append(p)
                       else: nps.append(p)
                  text = "\t".join(nps); filled_any = True

        if text != original_text: paragraph.text = text

    def _get_formatted_val(self, val: Any, placeholder: str, path: str) -> str:
        if isinstance(val, bool):
             if val: return " ☒" if "⎕" in placeholder else "[x]"
             else: return " ⎕" if "⎕" in placeholder else "[ ]"

        if val is None or val == "None" or str(val).strip() == "":
             if "⎕" in placeholder and len(placeholder) <= 2:
                  return " ⎕"
             text_len = len(placeholder) if not placeholder.startswith("⎕") else 1
             return " " * text_len

        sv = str(val)
        if "⎕" in placeholder:
             # Remove non-alphanumerics so things string cleanly into grids (e.g. 87654321-2-22 -> 87654321222)
             if any(k in path for k in ["tax_id", "tax_number", "id_number"]):
                  sv = "".join(c for c in sv if c.isalnum())
             num_boxes = placeholder.count("⎕")
             if len(sv) == 4 and num_boxes == 2 and sv.isdigit():
                  sv = sv[-2:]
             res = []; idx = 0
             for c in placeholder:
                  if c == "⎕": 
                       ch = sv[idx] if idx < len(sv) else " "
                       if num_boxes >= 10: res.append(f"{ch}  ") # Extra wide for ID fields
                       else: res.append(ch)
                       idx += 1
                  else: res.append(c)
             return "".join(res)
        
        # Standard text placeholder padding to prevent layout collapse
        v_str = f" {sv} "
        if len(v_str) < len(placeholder):
             return v_str + " " * (len(placeholder) - len(v_str))
        return v_str

    def _get_data_value(self, data: NormalizedData, path: str, context_label: str = "") -> Any:
        try:
            pts = path.split('.'); obj = data
            for pt in pts:
                if isinstance(obj, list):
                     if self.current_dep_idx < len(obj): obj = obj[self.current_dep_idx]
                     else: return None
                if isinstance(obj, dict):
                     obj = obj.get(pt)
                else: obj = getattr(obj, pt)
            if "birth_date" in path:
                pl = self._get_data_value_raw(data, path.replace("birth_date", "birth_place"))
                return f"{pl}, {obj}" if pl and obj else (obj or pl)
            if ("date" in path or "időpontja" in path) and context_label:
                ds = str(obj); m = re.search(r"(\d{4})[\.\s]+([^\s\.]+?)[\.\s]+(\d{1,2})", ds)
                if m:
                    if "év" in context_label: return m.group(1).strip()
                    if "hónap" in context_label: return m.group(2).strip()
                    if "nap" in context_label: return m.group(3).strip()
                p = [x.strip() for x in ds.replace('-', '.').split('.')]
                if len(p) >= 3:
                     if "év" in context_label: return p[0]
                     if "hónap" in context_label: return p[1]
                     if "nap" in context_label: return p[2]
            return obj
        except: return None

    def _get_data_value_raw(self, data: NormalizedData, path: str) -> Any:
        try:
            pts = path.split('.'); obj = data
            for pt in pts:
                if isinstance(obj, list):
                     if self.current_dep_idx < len(obj): obj = obj[self.current_dep_idx]
                     else: return None
                if isinstance(obj, dict):
                     obj = obj.get(pt)
                else: obj = getattr(obj, pt)
            return obj
        except: return None
